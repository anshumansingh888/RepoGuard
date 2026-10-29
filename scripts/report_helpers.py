import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F2F4F8"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="B0B8C4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, underline=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    if color_rgb:
        run.font.color.rgb = color_rgb

def add_p(doc, text="", bold_prefix="", italic=False, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, bold=True)
    if text:
        r_body = p.add_run(text)
        format_run(r_body, italic=italic)
    return p

def add_bullet(doc, bold_prefix, text, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    
    r_bullet = p.add_run("•  ")
    format_run(r_bullet, bold=True)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, bold=True)
    r_body = p.add_run(text)
    format_run(r_body)
    return p

def add_chapter_title(doc, number, title):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    
    r = p.add_run(f"CHAPTER {number}\n{title.upper()}")
    format_run(r, size_pt=16, bold=True, underline=True)
    return p

def add_section_title(doc, number_str, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    
    r = p.add_run(f"{number_str} {title}")
    format_run(r, size_pt=14, bold=True)
    return p

def add_subsection_title(doc, number_str, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    
    r = p.add_run(f"{number_str} {title}")
    format_run(r, size_pt=12, bold=True, italic=True)
    return p

def add_table_with_caption(doc, caption_text, headers, rows, col_widths=None):
    # Caption at TOP as per guidelines
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(12)
    p_cap.paragraph_format.space_after = Pt(4)
    r_cap = p_cap.add_run(caption_text)
    format_run(r_cap, size_pt=11, bold=True)
    
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "1F3A60")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            format_run(run, size_pt=10.5, bold=True, color_rgb=RGBColor(255, 255, 255))
            
    # Data rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                format_run(run, size_pt=10)
                
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(4)
    p_post.paragraph_format.space_after = Pt(6)

def add_figure_box(doc, caption_text, diagram_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F4F6F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="2B4C7E"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="2B4C7E"/>'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="2B4C7E"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="2B4C7E"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("\n".join(diagram_lines))
    format_run(r, font_name="Consolas", size_pt=9.5, color_rgb=RGBColor(20, 30, 45))
    
    # Caption BELOW figure as per guidelines
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption_text)
    format_run(r_cap, size_pt=11, bold=True, italic=True)

def add_code_snippet(doc, caption_text, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "0D1117")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(code_str)
    format_run(r, font_name="Consolas", size_pt=9, color_rgb=RGBColor(220, 230, 242))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    format_run(r_cap, size_pt=10.5, italic=True, bold=True)

def add_algorithm_box(doc, algo_num, title, inputs, outputs, steps):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(4)
    r_cap = p_cap.add_run(f"Algorithm {algo_num}: {title}")
    format_run(r_cap, size_pt=11.5, bold=True)
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F9FAFB")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="1F3A60"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1F3A60"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    
    r_in = p.add_run(f"Input: {inputs}\n")
    format_run(r_in, size_pt=10, bold=True)
    r_out = p.add_run(f"Output: {outputs}\n")
    format_run(r_out, size_pt=10, bold=True)
    
    r_div = p.add_run("―" * 55 + "\n")
    format_run(r_div, size_pt=9, color_rgb=RGBColor(140, 150, 165))
    
    for i, s in enumerate(steps, 1):
        r_step = p.add_run(f"{i:2d}:  {s}\n")
        format_run(r_step, font_name="Consolas", size_pt=9.5)
        
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(2)
    p_post.paragraph_format.space_after = Pt(8)

print("Report helper modules initialized.")

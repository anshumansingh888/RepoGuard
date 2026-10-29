import os
import sys
sys.path.insert(0, os.path.abspath("."))
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from scripts.report_helpers import (
    format_run, add_p, add_chapter_title, add_section_title, 
    set_table_borders, set_cell_shading, set_cell_margins
)
from scripts.report_content_ch1 import write_chapter_1
from scripts.report_content_ch2 import write_chapter_2
from scripts.report_content_ch3 import write_chapter_3
from scripts.report_content_ch4_ref import write_chapter_4_and_references

def update_front_matter(doc):
    print("Updating Front Matter...")
    
    # 1. Cover Page Title (P1)
    p_title = doc.paragraphs[1]
    p_title.text = "RepoGuard AI: Autonomous Multi-Agent Static Application Security Testing (SAST) and Automated Vulnerability Remediation Platform"
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p_title.runs:
        format_run(p_title.runs[0], size_pt=18, bold=True)

    # Clean P2
    doc.paragraphs[2].text = ""

    # Students (P14 to P19)
    doc.paragraphs[14].text = "Animesh Mudit (Roll No: ....................)"
    doc.paragraphs[15].text = "[Team Member 2 Name] (Roll No: ....................)"
    doc.paragraphs[16].text = "[Team Member 3 Name] (Roll No: ....................)"
    for idx in [17, 18, 19]:
        doc.paragraphs[idx].text = ""

    # Supervisor & Project ID (P22-P25)
    doc.paragraphs[23].text = "(Supervisor Name)"
    doc.paragraphs[24].text = "Assistant Professor"
    doc.paragraphs[25].text = "Project Id: PSIT-CSE-2026-RG01"

    # 2. Declaration (P34-P48)
    p_decl = doc.paragraphs[36]
    p_decl.text = (
        "This is to declare that the report entitled \"RepoGuard AI: Autonomous Multi-Agent Static Application "
        "Security Testing (SAST) and Automated Vulnerability Remediation Platform\", submitted by Animesh Mudit "
        "(Roll No: ....................), in partial fulfilment of the requirement for the award of the degree "
        "B.Tech. in Computer Science and Engineering at Pranveer Singh Institute of Technology, Kanpur, affiliated to "
        "Dr. A P J A K Technical University, Lucknow, comprises only our own work, and due acknowledgement has been "
        "made in the text to all other material used."
    )
    p_decl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if p_decl.runs:
        format_run(p_decl.runs[0], size_pt=12)

    doc.paragraphs[42].text = "Animesh Mudit (Roll No: ....................)"
    doc.paragraphs[43].text = "[Team Member 2 Name] (Roll No: ....................)"
    doc.paragraphs[44].text = "[Team Member 3 Name] (Roll No: ....................)"
    for idx in [45, 46, 47]:
        doc.paragraphs[idx].text = ""

    # 3. Certificate (P66-P70)
    p_cert = doc.paragraphs[68]
    p_cert.text = (
        "This is to certify that the report entitled \"RepoGuard AI: Autonomous Multi-Agent Static Application "
        "Security Testing (SAST) and Automated Vulnerability Remediation Platform\", which is submitted by "
        "Animesh Mudit (Roll No: ....................), in partial fulfilment of the requirements for the award "
        "of the degree B.Tech. in Computer Science & Engineering at Pranveer Singh Institute of Technology, Kanpur, "
        "affiliated to Dr. A P J A K Technical University, Lucknow, is a record of the candidates' own work carried out "
        "by them under my supervision. The matter embodied in this report is original and has not been submitted "
        "for the award of any other degree."
    )
    p_cert.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if p_cert.runs:
        format_run(p_cert.runs[0], size_pt=12)

    # 4. Acknowledgement (P93-P100)
    p_ack = doc.paragraphs[95]
    p_ack.text = (
        "It gives us great sense of pleasure to present our project undertaken during B.Tech. Third Year (Session: 2026-27). "
        "We owe a special debt of gratitude to our project supervisor, [Supervisor Name], Assistant Professor, Department "
        "of Computer Science and Engineering, Pranveer Singh Institute of Technology, Kanpur, for his constant support and "
        "guidance throughout the course of our work. His sincerity, thoroughness, and perseverance have been a constant "
        "source of inspiration for us."
    )
    p_ack.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if p_ack.runs:
        format_run(p_ack.runs[0], size_pt=12)

    # 5. Abstract (P101-P107)
    p_abs = doc.paragraphs[106]
    p_abs.text = (
        "In modern agile software development, static application security testing (SAST) tools are essential for "
        "early vulnerability discovery. However, conventional SAST solutions suffer from high false-positive rates "
        "and the 'Remediation Void'—they detect security flaws but place the burden of patch research, synthesis, "
        "and verification entirely on human developers. Furthermore, sending proprietary source code to cloud-hosted "
        "language models introduces grave data exfiltration risks. This report presents RepoGuard AI, an autonomous, "
        "multi-agent SAST and automated remediation platform driven by a localized large language model (Qwen 2.5 Coder 7B) "
        "executing privately via the Ollama runtime. RepoGuard decouples static analysis into two specialized agents: "
        "Agent 1 (The Auditor), which executes Abstract Syntax Tree (AST) taint analysis constrained by eight ground-truth "
        "false-positive exclusion rules and type-safe Pydantic JSON contracts; and Agent 2 (The Fixer), which performs surgical "
        "code remediation using strictly sanctioned Python standard library primitives. An autonomous Gatekeeper supervision "
        "engine implements an iterative convergence feedback loop that automatically re-audits remediated code up to three "
        "passes, guaranteeing mathematically verified zero-residual vulnerability states while executing a single-pass Fast Exit "
        "for pre-secured code. Experimental evaluation across standard benchmarks demonstrates that RepoGuard achieves 100% "
        "precision and recall across six critical vulnerability classes (SQL Injection, OS Command Injection, Path Traversal, "
        "Hardcoded Secrets, Weak Cryptography, and Mutable Default Arguments), eliminating 100% of residual vulnerabilities "
        "within two convergence passes. A reactive Streamlit dashboard delivers real-time telemetry, side-by-side comparative "
        "inspection, and unified Git diff patch export, establishing a production-grade, zero-cost blueprint for autonomous DevSecOps."
    )
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if p_abs.runs:
        format_run(p_abs.runs[0], size_pt=12)

    # 6. Update Table of Contents (Table 2)
    tbl_toc = doc.tables[2]
    # Clear existing rows except header
    for _ in range(len(tbl_toc.rows) - 1):
        tr = tbl_toc.rows[-1]._tr
        tbl_toc._tbl.remove(tr)

    toc_data = [
        ("1", "Declaration", "ii"),
        ("2", "Certificate", "iii"),
        ("3", "Acknowledgement", "iv"),
        ("4", "Abstract", "v"),
        ("5", "List of Figures", "vi"),
        ("6", "List of Tables", "vii"),
        ("7", "Chapter 1: Introduction", "1"),
        ("", "    1.1 Background and Overview of Application Security", "1"),
        ("", "    1.2 Evolution of Application Security Testing Paradigms", "3"),
        ("", "    1.3 The Shift-Left Paradigm and DevSecOps Integration", "5"),
        ("", "    1.4 Problem Statement: The Remediation Void in Modern SAST", "6"),
        ("", "    1.5 Critical Analysis of Existing Tools and State-of-the-Art", "8"),
        ("", "    1.6 Security and Privacy Risks of Cloud-Hosted Code LLMs", "9"),
        ("", "    1.7 Research Motivation and Architectural Rationale", "10"),
        ("", "    1.8 Specific Project Objectives", "10"),
        ("", "    1.9 Scope and Assumptions", "11"),
        ("", "    1.10 Organization of the Report", "11"),
        ("8", "Chapter 2: Proposed Methodology", "12"),
        ("", "    2.1 Architectural Blueprint of the RepoGuard Framework", "12"),
        ("", "    2.2 Theoretical Foundations of Static Taint Analysis", "14"),
        ("", "    2.3 Agent 1 (The Security Auditor): Design and Prompt Engineering", "16"),
        ("", "    2.4 Strict Pydantic Data Contracts and Validated JSON Schemas", "17"),
        ("", "    2.5 Deterministic Sampling Parameters", "18"),
        ("", "    2.6 Ground-Truth Exclusion Rules (False-Positive Suppression)", "19"),
        ("", "    2.7 The Autonomous Gatekeeper Decision Mechanism", "21"),
        ("", "    2.8 Fast-Path Zero-Vulnerability Bypass Logic", "22"),
        ("", "    2.9 Agent 2 (The Secure Fixer): Surgical Patch Synthesis", "22"),
        ("", "    2.10 Standard Library Security Guardrails", "23"),
        ("", "    2.11 Docstring and Comment Sanitization Protocol", "25"),
        ("", "    2.12 Post-Processing Engine and Mandatory Import Injection", "25"),
        ("", "    2.13 Closed-Loop Self-Healing Convergence Mechanism", "26"),
        ("", "    2.14 Algorithmic Formulations (Algorithm 1 & 2)", "26"),
        ("", "    2.15 Interactive Streamlit Web Interface and Telemetry", "27"),
        ("9", "Chapter 3: Results and Discussions", "28"),
        ("", "    3.1 Experimental Setup and Testbed Environment", "28"),
        ("", "    3.2 Benchmark Dataset and Target Vulnerability Profiles", "29"),
        ("", "    3.3 Quantitative Evaluation Metrics", "30"),
        ("", "    3.4 In-Depth Empirical Case Studies", "31"),
        ("", "    3.5 Convergence Loop Trajectory and Flaw Elimination", "34"),
        ("", "    3.6 Comparative Performance Analysis: RepoGuard vs Existing Tools", "36"),
        ("", "    3.7 Ablation Studies", "37"),
        ("", "    3.8 User Experience and Telemetry Evaluation", "38"),
        ("", "    3.9 Threats to Validity", "39"),
        ("10", "Chapter 4: Conclusion and Future Scope", "40"),
        ("", "    4.1 Summary of Contributions", "40"),
        ("", "    4.2 Engineering and Pedagogical Learnings", "41"),
        ("", "    4.3 Practical Limitations of Current Implementation", "42"),
        ("", "    4.4 Roadmap for Future Enhancements (Major Project Scope)", "42"),
        ("", "    4.5 Concluding Remarks", "44"),
        ("11", "References", "45")
    ]

    for s_no, title, page in toc_data:
        row = tbl_toc.add_row()
        cells = row.cells
        cells[0].text = s_no
        cells[1].text = title
        cells[2].text = page
        set_cell_margins(cells[0], top=40, bottom=40, left=80, right=80)
        set_cell_margins(cells[1], top=40, bottom=40, left=80, right=80)
        set_cell_margins(cells[2], top=40, bottom=40, left=80, right=80)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    format_run(run, size_pt=10.5)

    print("Front Matter successfully updated.")

def remove_template_guidelines(doc):
    print("Removing placeholder guidelines...")
    body = doc.element.body
    # Elements 173 to 198 were guideline paragraphs
    to_remove = []
    found_guidelines = False
    for child in list(body):
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            text = child.text.strip()
            if text.startswith('Content Guidelines:'):
                found_guidelines = True
            if found_guidelines:
                to_remove.append(child)
                
    for elem in to_remove:
        body.remove(elem)
    print(f"Removed {len(to_remove)} placeholder elements.")

def add_list_of_figures_and_tables(doc):
    # List of Figures (Page vi)
    doc.add_page_break()
    p_fig_head = doc.add_paragraph()
    p_fig_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fig_head.paragraph_format.space_before = Pt(18)
    p_fig_head.paragraph_format.space_after = Pt(14)
    r_fig = p_fig_head.add_run("LIST OF FIGURES")
    format_run(r_fig, size_pt=16, bold=True, underline=True)

    fig_headers = ["Figure No.", "Figure Title", "Page No."]
    fig_rows = [
        ["Figure 2.1", "Architectural Blueprint and Execution Pipeline of RepoGuard Framework", "13"],
        ["Figure 2.2", "AST Taint Flow Control Graph from Sources to Sinks", "15"],
        ["Figure 2.3", "Gatekeeper Closed-Loop Convergence State Machine", "20"],
        ["Figure 3.1", "Iterative Vulnerability Convergence Trajectory Across Passes", "35"],
        ["Figure 3.2", "Streamlit Dashboard Cyberpunk Interface and Unified Diff Viewer", "38"]
    ]
    
    tbl_fig = doc.add_table(rows=len(fig_rows) + 1, cols=3)
    tbl_fig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_fig)
    for i, h in enumerate(fig_headers):
        tbl_fig.rows[0].cells[i].text = h
        set_cell_shading(tbl_fig.rows[0].cells[i], "1F3A60")
        set_cell_margins(tbl_fig.rows[0].cells[i], top=100, bottom=100, left=120, right=120)
        p = tbl_fig.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            format_run(r, size_pt=10.5, bold=True, color_rgb=RGBColor(255, 255, 255))
            
    for idx, r_data in enumerate(fig_rows):
        cells = tbl_fig.rows[idx + 1].cells
        for c_idx, val in enumerate(r_data):
            cells[c_idx].text = val
            set_cell_margins(cells[c_idx], top=60, bottom=60, left=100, right=100)
            p = cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 1 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                format_run(r, size_pt=10)

    # List of Tables (Page vii)
    doc.add_page_break()
    p_tbl_head = doc.add_paragraph()
    p_tbl_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbl_head.paragraph_format.space_before = Pt(18)
    p_tbl_head.paragraph_format.space_after = Pt(14)
    r_tbl = p_tbl_head.add_run("LIST OF TABLES")
    format_run(r_tbl, size_pt=16, bold=True, underline=True)

    tbl_headers = ["Table No.", "Table Title", "Page No."]
    tbl_rows = [
        ["Table 1.1", "Comparative Taxonomy of Application Security Testing (AST) Paradigms", "4"],
        ["Table 2.1", "Formal Ground-Truth False-Positive Exclusion Rules in Agent 1", "19"],
        ["Table 3.1", "Experimental Hardware and Software Testbed Specifications", "28"],
        ["Table 3.2", "Step-by-Step Remediation Mapping on Composite Benchmark Sample", "32"],
        ["Table 3.3", "Iterative Convergence Telemetry Across Benchmark Targets", "34"],
        ["Table 3.4", "Comparative Performance Matrix: RepoGuard vs Existing SAST and AI Solutions", "36"]
    ]
    
    tbl_t = doc.add_table(rows=len(tbl_rows) + 1, cols=3)
    tbl_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_t)
    for i, h in enumerate(tbl_headers):
        tbl_t.rows[0].cells[i].text = h
        set_cell_shading(tbl_t.rows[0].cells[i], "1F3A60")
        set_cell_margins(tbl_t.rows[0].cells[i], top=100, bottom=100, left=120, right=120)
        p = tbl_t.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            format_run(r, size_pt=10.5, bold=True, color_rgb=RGBColor(255, 255, 255))
            
    for idx, r_data in enumerate(tbl_rows):
        cells = tbl_t.rows[idx + 1].cells
        for c_idx, val in enumerate(r_data):
            cells[c_idx].text = val
            set_cell_margins(cells[c_idx], top=60, bottom=60, left=100, right=100)
            p = cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 1 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                format_run(r, size_pt=10)

def main():
    template_path = "docs/Third Year Mini Project Report 26-27.docx"
    output_path = "docs/Third Year Mini Project Report 26-27.docx"
    backup_path = "docs/RepoGuard_AKTU_Mini_Project_Report.docx"

    print(f"Loading template from '{template_path}'...")
    doc = docx.Document(template_path)

    # 1. Update Front Matter
    update_front_matter(doc)

    # 2. Remove placeholder guidelines
    remove_template_guidelines(doc)

    # 3. Add List of Figures and List of Tables
    add_list_of_figures_and_tables(doc)

    # 4. Generate Chapters
    print("Writing Chapter 1: Introduction...")
    write_chapter_1(doc)

    print("Writing Chapter 2: Proposed Methodology...")
    write_chapter_2(doc)

    print("Writing Chapter 3: Results and Discussions...")
    write_chapter_3(doc)

    print("Writing Chapter 4: Conclusion & Future Scope and References...")
    write_chapter_4_and_references(doc)

    # 5. Save Document
    print(f"Saving final report to '{output_path}'...")
    doc.save(output_path)
    print(f"Saving backup copy to '{backup_path}'...")
    doc.save(backup_path)
    print("[SUCCESS] All operations completed successfully!")

if __name__ == "__main__":
    main()
